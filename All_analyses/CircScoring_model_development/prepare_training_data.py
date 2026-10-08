#!/usr/bin/env python3
"""Extract the exact P1/N1 CircScoring development set from Excel."""

from __future__ import annotations

import argparse
import csv
import hashlib
from pathlib import Path
import sys

try:
    import openpyxl
except ImportError as error:
    raise SystemExit(
        "openpyxl is required. Install version 3.1.5 with: "
        "python -m pip install openpyxl==3.1.5"
    ) from error


EXPECTED_FEATURES = [
    "donor_site_at_the_annotated_boundary",
    "acceptor_site_at_the_annotated_boundary",
    "donor_acceptor_sites_at_the_same_transcript_isoform",
    "has_AS_event(donor)",
    "has_AS_event(acceptor)",
    "phyloP(acceptor_in)",
    "phyloP(acceptor_out)",
    "phyloP(donor_in)",
    "phyloP(donor_out)",
    "phastCons(acceptor_in)",
    "phastCons(acceptor_out)",
    "phastCons(donor_in)",
    "phastCons(donor_out)",
    "MAXENT(donor)",
    "MAXENT(acceptor)",
    "7 DB",
    "FL-circAS or circFL_seq",
    "mouse_conserved",
]


def format_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return format(value, ".17g")
    return str(value)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("source_xlsx", type=Path)
    parser.add_argument("output_tsv", type=Path)
    parser.add_argument("--sheet", default="6cols")
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(
        args.source_xlsx, read_only=True, data_only=True
    )
    if args.sheet not in workbook.sheetnames:
        raise ValueError(f"Worksheet not found: {args.sheet}")
    worksheet = workbook[args.sheet]
    iterator = worksheet.iter_rows(values_only=True)
    source_header = [str(value) if value is not None else "" for value in next(iterator)]
    required = ["circRNA_id", "P1", "N1", *EXPECTED_FEATURES]
    missing = [name for name in required if name not in source_header]
    if missing:
        raise ValueError("Missing required column(s): " + ", ".join(missing))
    index = {name: source_header.index(name) for name in required}

    output_rows: list[list[object]] = []
    seen_ids: set[str] = set()
    positive = 0
    negative = 0
    for excel_row_number, row in enumerate(iterator, start=2):
        p1 = row[index["P1"]]
        n1 = row[index["N1"]]
        if not ((p1 == 1 and n1 == 0) or (p1 == 0 and n1 == 1)):
            continue
        circ_id = str(row[index["circRNA_id"]]).strip()
        if not circ_id:
            raise ValueError(f"Missing circRNA_id at Excel row {excel_row_number}")
        if circ_id in seen_ids:
            raise ValueError(f"Duplicate P1/N1 circRNA_id: {circ_id}")
        seen_ids.add(circ_id)
        values = [row[index[name]] for name in required]
        for name, value in zip(required[3:], values[3:]):
            if value is None or (isinstance(value, str) and not value.strip()):
                raise ValueError(
                    f"Missing {name!r} for {circ_id} at Excel row {excel_row_number}"
                )
            try:
                float(value)
            except (TypeError, ValueError) as error:
                raise ValueError(
                    f"Nonnumeric {name!r} for {circ_id}: {value!r}"
                ) from error
        output_rows.append(values)
        positive += int(p1 == 1)
        negative += int(n1 == 1)

    workbook.close()
    if (len(output_rows), positive, negative) != (1301, 797, 504):
        raise ValueError(
            "Unexpected P1/N1 population: "
            f"rows={len(output_rows)}, P1={positive}, N1={negative}"
        )

    args.output_tsv.parent.mkdir(parents=True, exist_ok=True)
    with args.output_tsv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(required)
        writer.writerows([[format_value(value) for value in row] for row in output_rows])

    print(f"Saved {len(output_rows)} rows to {args.output_tsv}")
    print(f"P1={positive}; N1={negative}")
    print(f"SHA-256={sha256(args.output_tsv)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())

