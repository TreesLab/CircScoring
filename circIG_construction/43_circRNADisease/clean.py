#!/usr/bin/env python3
"""Clean circRNADisease entries and extract normalized circRNA IDs."""

from __future__ import annotations

import argparse
import csv
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_RAW = BASE_DIR / "raw" / "circRNADisease_V3_circrna_details.xlsx"
DEFAULT_CIRCBASE_ID_MAP = BASE_DIR / "resources" / "circBase_id_updates.tsv"
DEFAULT_CLEAN_DIR = BASE_DIR / "clean"

ENTRY_OUTPUT = "43_circRNADisease.entries.clean.tsv"
ID_OUTPUT = "43_circRNADisease.circRNA_ids.clean.tsv"
SUMMARY_OUTPUT = "43_circRNADisease.clean.summary.tsv"
ID_MAP_COLUMNS = ["old_circRNA_id", "new_circRNA_id", "status", "search_result_count"]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw", type=Path, default=DEFAULT_RAW, help=f"Raw Excel file. Default: {DEFAULT_RAW}")
    parser.add_argument(
        "--circbase-id-map",
        type=Path,
        default=DEFAULT_CIRCBASE_ID_MAP,
        help=f"Versioned old-to-current circBase ID mapping. Default: {DEFAULT_CIRCBASE_ID_MAP}",
    )
    parser.add_argument("--clean-dir", type=Path, default=DEFAULT_CLEAN_DIR, help=f"Clean output directory. Default: {DEFAULT_CLEAN_DIR}")
    parser.add_argument("--force", action="store_true", help="Overwrite existing outputs.")
    return parser.parse_args()


def resolve(path: Path) -> Path:
    return path.expanduser().resolve()


def normalize(value: object) -> str:
    if value is None:
        return "NA"
    text = str(value).strip()
    if not text or text.lower() == "nan":
        return "NA"
    return " ".join(text.split())


def refuse_overwrite(paths: list[Path], force: bool) -> None:
    if force:
        return
    existing = [str(path) for path in paths if path.exists()]
    if existing:
        raise FileExistsError("Output exists; use --force to overwrite: " + ", ".join(existing))


def load_id_map(path: Path) -> dict[str, str]:
    mapping: dict[str, str] = {}
    with path.open(newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        missing = set(ID_MAP_COLUMNS) - set(reader.fieldnames or [])
        if missing:
            raise ValueError(f"{path} missing required columns: {', '.join(sorted(missing))}")
        for row in reader:
            if row["status"] != "unique_match":
                continue
            old_id = normalize(row["old_circRNA_id"])
            new_id = normalize(row["new_circRNA_id"])
            if "NA" in {old_id, new_id}:
                raise ValueError(f"{path} contains an invalid unique_match row: {row}")
            previous = mapping.get(old_id)
            if previous is not None and previous != new_id:
                raise ValueError(f"{path} maps {old_id} to both {previous} and {new_id}")
            mapping[old_id] = new_id
    return mapping


def read_excel(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    try:
        header_values = next(rows)
    except StopIteration as exc:
        raise ValueError(f"{path} is empty") from exc

    headers = [normalize(value) for value in header_values]
    if "circrna_id" not in headers:
        raise ValueError(f"{path} missing required column: circrna_id")

    output: list[dict[str, str]] = []
    for values in rows:
        row = {header: normalize(value) for header, value in zip(headers, values) if header != "NA"}
        if row and any(value != "NA" for value in row.values()):
            output.append(row)
    workbook.close()
    return headers, output


def write_tsv(path: Path, columns: list[str], rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=columns,
            delimiter="\t",
            lineterminator="\n",
            extrasaction="ignore",
        )
        writer.writeheader()
        writer.writerows(rows)


def clean(raw_path: Path, id_map_path: Path, clean_dir: Path, force: bool) -> Counter[str]:
    entries_output = clean_dir / ENTRY_OUTPUT
    ids_output = clean_dir / ID_OUTPUT
    summary_output = clean_dir / SUMMARY_OUTPUT
    refuse_overwrite([entries_output, ids_output, summary_output], force)

    id_map = load_id_map(id_map_path)
    raw_headers, raw_rows = read_excel(raw_path)
    counts: Counter[str] = Counter(raw_rows=len(raw_rows))
    entries: list[dict[str, str]] = []
    circ_ids: set[str] = set()
    updated_ids: set[str] = set()

    for raw_row in raw_rows:
        original_id = normalize(raw_row.get("circrna_id"))
        if original_id == "NA":
            counts["rows_without_circRNA_id"] += 1
            continue

        counts["rows_with_circRNA_id"] += 1
        circ_id = id_map.get(original_id, original_id)
        if circ_id != original_id:
            counts["rows_with_updated_circRNA_id"] += 1
            updated_ids.add(original_id)
        entries.append({"circRNA_id": circ_id, **raw_row})
        circ_ids.add(circ_id)

    counts["clean_entry_rows"] = len(entries)
    counts["clean_unique_circRNA_ids"] = len(circ_ids)
    counts["updated_unique_circRNA_ids"] = len(updated_ids)

    write_tsv(entries_output, ["circRNA_id", *raw_headers], entries)
    write_tsv(ids_output, ["circRNA_id"], [{"circRNA_id": circ_id} for circ_id in sorted(circ_ids)])
    write_tsv(
        summary_output,
        ["metric", "value"],
        [{"metric": key, "value": str(value)} for key, value in counts.items()],
    )
    return counts


def main() -> int:
    args = parse_args()
    counts = clean(resolve(args.raw), resolve(args.circbase_id_map), resolve(args.clean_dir), args.force)
    for key, value in counts.items():
        print(f"{key}\t{value}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
