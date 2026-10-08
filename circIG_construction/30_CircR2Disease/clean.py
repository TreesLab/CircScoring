#!/usr/bin/env python3
"""Clean CircR2Disease entries into preprocessed and circRNA coordinate tables."""

from __future__ import annotations

import argparse
import csv
import re
import unicodedata
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_RAW = BASE_DIR / "raw" / "The circRNA-disease entries.xlsx"
DEFAULT_TECHNIQUE_MAPPING = BASE_DIR / "resources" / "experimental_technique_name_mapping.tsv"
DEFAULT_PREPROCESSED_OUTPUT = BASE_DIR / "clean" / "30_CircR2Disease.entries.human.preprocessed.tsv"
DEFAULT_CLEAN_OUTPUT = BASE_DIR / "clean" / "30_CircR2Disease.circRNAs.clean.tsv"
DEFAULT_AUDIT_OUTPUT = BASE_DIR / "clean" / "30_CircR2Disease.clean.audit.tsv"
DEFAULT_TECHNIQUE_AUDIT_OUTPUT = BASE_DIR / "clean" / "30_CircR2Disease.experimental_techniques.unmapped.tsv"
DEFAULT_SUMMARY_OUTPUT = BASE_DIR / "clean" / "30_CircR2Disease.clean.summary.tsv"

PREPROCESSED_COLUMNS = [
    "CRD ID",
    "circRNA Name",
    "Synonyms",
    "Gene Symbol",
    "Disease Name",
    "Expression pattern",
    "PubMed ID",
    "Region",
    "Strand",
    "Species",
    "Experimental techniques",
    "Brief description",
    "Title",
    "circRNA_id",
    "region_normalized",
    "strand_normalized",
    "experimental_techniques_normalized",
]
CLEAN_COLUMNS = ["circRNA_id", "chrom", "start", "end", "strand", "gene_symbol"]
AUDIT_COLUMNS = ["source", "reason", "circRNA_id", "region", "strand", "gene_symbol", "CRD ID"]
TECHNIQUE_AUDIT_COLUMNS = ["token", "count"]
SUMMARY_KEYS = [
    "raw_rows",
    "human_rows",
    "preprocessed_rows",
    "clean_rows_before_uniq",
    "duplicate_full_rows_removed",
    "clean_rows",
    "audit_rows",
    "unmapped_experimental_technique_tokens",
    "unique_circRNA_ids",
    "unique_gene_symbols",
    "gene_symbol_na_rows",
    "strand:+",
    "strand:-",
    "strand:NA",
]

MISSING_VALUES = {"", "NA", "N/A", "NULL", "NONE", "NAN"}
DASH_CHARS = {
    "\u2010",
    "\u2011",
    "\u2012",
    "\u2013",
    "\u2014",
    "\u2015",
    "\u2212",
    "\ufe58",
    "\ufe63",
    "\uff0d",
}
SPACE_CHARS = {
    "\u00a0",
    "\u1680",
    "\u2000",
    "\u2001",
    "\u2002",
    "\u2003",
    "\u2004",
    "\u2005",
    "\u2006",
    "\u2007",
    "\u2008",
    "\u2009",
    "\u200a",
    "\u202f",
    "\u205f",
    "\u3000",
}
ZERO_WIDTH_CHARS = {
    "\u200b",
    "\u200c",
    "\u200d",
    "\ufeff",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Clean CircR2Disease circRNA-disease entries.")
    parser.add_argument("--raw", type=Path, default=DEFAULT_RAW, help=f"Raw xlsx file. Default: {DEFAULT_RAW}")
    parser.add_argument(
        "--technique-mapping",
        type=Path,
        default=DEFAULT_TECHNIQUE_MAPPING,
        help=f"Experimental technique mapping TSV. Default: {DEFAULT_TECHNIQUE_MAPPING}",
    )
    parser.add_argument(
        "--preprocessed-output",
        type=Path,
        default=DEFAULT_PREPROCESSED_OUTPUT,
        help=f"Human preprocessed entries TSV. Default: {DEFAULT_PREPROCESSED_OUTPUT}",
    )
    parser.add_argument("--output", type=Path, default=DEFAULT_CLEAN_OUTPUT, help=f"Clean output TSV. Default: {DEFAULT_CLEAN_OUTPUT}")
    parser.add_argument(
        "--audit-output",
        type=Path,
        default=DEFAULT_AUDIT_OUTPUT,
        help=f"Clean audit output TSV. Default: {DEFAULT_AUDIT_OUTPUT}",
    )
    parser.add_argument(
        "--technique-audit-output",
        type=Path,
        default=DEFAULT_TECHNIQUE_AUDIT_OUTPUT,
        help=f"Unmapped technique audit TSV. Default: {DEFAULT_TECHNIQUE_AUDIT_OUTPUT}",
    )
    parser.add_argument(
        "--summary-output",
        type=Path,
        default=DEFAULT_SUMMARY_OUTPUT,
        help=f"Summary output TSV. Default: {DEFAULT_SUMMARY_OUTPUT}",
    )
    parser.add_argument("--force", action="store_true", help="Overwrite existing outputs.")
    return parser.parse_args()


def resolve(path: Path) -> Path:
    return path.expanduser().resolve()


def refuse_overwrite(paths: list[Path], force: bool) -> None:
    if force:
        return
    existing = [path for path in paths if path.exists()]
    if existing:
        raise FileExistsError(f"Output exists; use --force to overwrite: {', '.join(str(path) for path in existing)}")


def normalize_text(value: object) -> str:
    if value is None:
        return "NA"
    text = str(value)
    text = unicodedata.normalize("NFKC", text)
    for char in DASH_CHARS:
        text = text.replace(char, "-")
    for char in SPACE_CHARS:
        text = text.replace(char, " ")
    for char in ZERO_WIDTH_CHARS:
        text = text.replace(char, "")
    text = text.replace("\t", " ")
    text = re.sub(r"\s+", " ", text).strip()
    if text.upper() in MISSING_VALUES:
        return "NA"
    return text


def normalize_circ_id(value: object) -> str:
    text = normalize_text(value)
    if text == "NA":
        return "NA"
    text = text.replace(" :", ":")
    text = text.replace(": ", ":")
    return text.strip()


def normalize_region(value: object) -> str:
    text = normalize_text(value)
    if text == "NA":
        return "NA"
    text = text.replace(",", "")
    text = re.sub(r"\s+", "", text)
    return text


def normalize_strand(value: object) -> str:
    text = normalize_text(value)
    if text in {"+", "-"}:
        return text
    return "NA"


def normalize_gene_symbol(value: object) -> str:
    text = normalize_text(value)
    return text if text != "NA" else "NA"


def normalize_technique_token(value: object) -> str:
    text = normalize_text(value)
    if text == "NA":
        return "NA"
    text = text.replace("_x0002_", "")
    text = re.sub(r"\betc\.?$", "", text, flags=re.IGNORECASE).strip(" ,;")
    text = re.sub(r"\s*/\s*", "/", text)
    return normalize_text(text)


def load_technique_mapping(path: Path) -> dict[str, str]:
    mapping: dict[str, str] = {}
    with path.open(newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        required = {"Corrected experimental technique", "Original values"}
        missing = required - set(reader.fieldnames or [])
        if missing:
            raise ValueError(f"{path} missing required columns: {', '.join(sorted(missing))}")
        for row in reader:
            corrected = normalize_technique_token(row["Corrected experimental technique"])
            originals = row["Original values"]
            for original in originals.split(","):
                token = normalize_technique_token(original)
                if token != "NA":
                    mapping[token.lower()] = corrected
            if corrected != "NA":
                mapping[corrected.lower()] = corrected
    return mapping


def normalize_techniques(value: object, mapping: dict[str, str], unmapped: Counter[str]) -> str:
    text = normalize_text(value)
    if text == "NA":
        return "NA"
    text = re.sub(r"\betc\.?", "", text, flags=re.IGNORECASE)
    tokens = [normalize_technique_token(token) for token in re.split(r"[,;]", text)]
    normalized: list[str] = []
    seen: set[str] = set()
    for token in tokens:
        if token == "NA":
            continue
        corrected = mapping.get(token.lower())
        if corrected is None:
            corrected = token
            unmapped[token] += 1
        if corrected not in seen:
            normalized.append(corrected)
            seen.add(corrected)
    return ",".join(normalized) if normalized else "NA"


def read_xlsx_rows(path: Path) -> tuple[list[str], list[dict[str, object]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    header_values = next(rows)
    header = [normalize_text(value) for value in header_values]
    output_rows: list[dict[str, object]] = []
    for values in rows:
        row = {header[index]: values[index] if index < len(values) else None for index in range(len(header))}
        output_rows.append(row)
    workbook.close()
    return header, output_rows


def chrom_key(chrom: str) -> tuple[int, str]:
    value = chrom.removeprefix("chr")
    if value.isdigit():
        return (int(value), "")
    special = {"X": 23, "Y": 24, "M": 25, "MT": 25}
    return (special.get(value, 1000), value)


def parse_region(region: str) -> tuple[str, int, int]:
    match = re.fullmatch(r"(chr[0-9A-Za-z]+):([0-9]+)-([0-9]+)", region)
    if not match:
        raise ValueError("unparseable_region")
    chrom, start, end = match.groups()
    start_int = int(start)
    end_int = int(end)
    if start_int < 0 or end_int < 0:
        raise ValueError("negative_coordinate")
    if start_int > end_int:
        raise ValueError("start_greater_than_end")
    return chrom, start_int, end_int


def write_dict_rows(path: Path, columns: list[str], rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def write_summary(path: Path, counts: Counter[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(["metric", "value"])
        for key in SUMMARY_KEYS:
            writer.writerow([key, counts[key]])


def main() -> int:
    args = parse_args()
    raw = resolve(args.raw)
    technique_mapping = resolve(args.technique_mapping)
    preprocessed_output = resolve(args.preprocessed_output)
    clean_output = resolve(args.output)
    audit_output = resolve(args.audit_output)
    technique_audit_output = resolve(args.technique_audit_output)
    summary_output = resolve(args.summary_output)

    refuse_overwrite([preprocessed_output, clean_output, audit_output, technique_audit_output, summary_output], args.force)

    mapping = load_technique_mapping(technique_mapping)
    header, raw_rows = read_xlsx_rows(raw)
    required = set(PREPROCESSED_COLUMNS[:13])
    missing = required - set(header)
    if missing:
        raise ValueError(f"{raw} missing required columns: {', '.join(sorted(missing))}")

    counts: Counter[str] = Counter()
    unmapped_techniques: Counter[str] = Counter()
    preprocessed_rows: list[dict[str, str]] = []

    counts["raw_rows"] = len(raw_rows)
    for row in raw_rows:
        species = normalize_text(row.get("Species"))
        if species.lower() != "human":
            continue
        counts["human_rows"] += 1

        normalized_row = {column: normalize_text(row.get(column)) for column in PREPROCESSED_COLUMNS[:13]}
        normalized_row["Species"] = "Human"
        normalized_row["circRNA_id"] = normalize_circ_id(row.get("circRNA Name"))
        normalized_row["region_normalized"] = normalize_region(row.get("Region"))
        normalized_row["strand_normalized"] = normalize_strand(row.get("Strand"))
        normalized_row["experimental_techniques_normalized"] = normalize_techniques(
            row.get("Experimental techniques"), mapping, unmapped_techniques
        )
        preprocessed_rows.append(normalized_row)

    counts["preprocessed_rows"] = len(preprocessed_rows)
    write_dict_rows(preprocessed_output, PREPROCESSED_COLUMNS, preprocessed_rows)

    clean_rows: list[tuple[str, str, str, str, str, str]] = []
    audit_rows: list[dict[str, str]] = []
    seen: set[tuple[str, str, str, str, str, str]] = set()
    circ_ids: set[str] = set()
    gene_symbols: set[str] = set()

    for row in preprocessed_rows:
        circ_id = row["circRNA_id"]
        region = row["region_normalized"]
        strand = row["strand_normalized"]
        gene_symbol = normalize_gene_symbol(row["Gene Symbol"])
        try:
            if circ_id == "NA":
                raise ValueError("missing_circRNA_id")
            chrom, start, end = parse_region(region)
        except ValueError as exc:
            audit_rows.append(
                {
                    "source": "preprocessed_entries",
                    "reason": str(exc),
                    "circRNA_id": circ_id,
                    "region": region,
                    "strand": strand,
                    "gene_symbol": gene_symbol,
                    "CRD ID": row["CRD ID"],
                }
            )
            continue

        clean_row = (circ_id, chrom, str(start), str(end), strand, gene_symbol)
        counts["clean_rows_before_uniq"] += 1
        if clean_row in seen:
            counts["duplicate_full_rows_removed"] += 1
            continue
        seen.add(clean_row)
        clean_rows.append(clean_row)
        circ_ids.add(circ_id)
        if gene_symbol == "NA":
            counts["gene_symbol_na_rows"] += 1
        else:
            gene_symbols.add(gene_symbol)
        counts[f"strand:{strand}"] += 1

    clean_rows.sort(key=lambda row: (chrom_key(row[1]), int(row[2]), int(row[3]), row[4], row[0]))
    clean_dict_rows = [
        {
            "circRNA_id": row[0],
            "chrom": row[1],
            "start": row[2],
            "end": row[3],
            "strand": row[4],
            "gene_symbol": row[5],
        }
        for row in clean_rows
    ]
    write_dict_rows(clean_output, CLEAN_COLUMNS, clean_dict_rows)
    write_dict_rows(audit_output, AUDIT_COLUMNS, audit_rows)

    technique_audit_rows = [{"token": token, "count": str(count)} for token, count in sorted(unmapped_techniques.items())]
    write_dict_rows(technique_audit_output, TECHNIQUE_AUDIT_COLUMNS, technique_audit_rows)

    counts["clean_rows"] = len(clean_rows)
    counts["audit_rows"] = len(audit_rows)
    counts["unmapped_experimental_technique_tokens"] = len(unmapped_techniques)
    counts["unique_circRNA_ids"] = len(circ_ids)
    counts["unique_gene_symbols"] = len(gene_symbols)
    write_summary(summary_output, counts)

    print(f"preprocessed_rows\t{counts['preprocessed_rows']}")
    print(f"clean_rows\t{counts['clean_rows']}")
    print(f"audit_rows\t{counts['audit_rows']}")
    print(f"unmapped_experimental_technique_tokens\t{counts['unmapped_experimental_technique_tokens']}")
    print(f"preprocessed_output\t{preprocessed_output}")
    print(f"clean_output\t{clean_output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
