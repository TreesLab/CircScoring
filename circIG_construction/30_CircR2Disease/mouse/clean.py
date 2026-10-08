#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import gzip
import json
import re
import sys
import unicodedata
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

try:
    from openpyxl import load_workbook
except ImportError:  # only needed by selected databases
    load_workbook = None

csv.field_size_limit(sys.maxsize)
OUTPUT_COLUMNS = ["circRNA_id", "chrom", "start", "end", "strand", "gene_symbol"]
NULL_VALUES = {"", "-", ".", "*", "/", "NA", "N/A", "NULL", "NONE", "NAN", "None", "none", "null", "Unknown", "unknown"}

SCRIPT_DIR = Path(__file__).resolve().parent
DATABASE = SCRIPT_DIR.parent.name
DEFAULT_RAW_DIR = SCRIPT_DIR / "raw"
DEFAULT_OUTPUT = SCRIPT_DIR / "clean" / f"{DATABASE}.circRNAs.clean.tsv"
DEFAULT_INVALID_OUTPUT = SCRIPT_DIR / "clean" / f"{DATABASE}.circRNAs.clean.invalid.tsv"
DEFAULT_SUMMARY_OUTPUT = SCRIPT_DIR / "clean" / f"{DATABASE}.circRNAs.clean.summary.tsv"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=f"Clean {DATABASE} mouse circRNA raw data into six columns.")
    parser.add_argument("--raw-dir", type=Path, default=DEFAULT_RAW_DIR)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--invalid-output", type=Path, default=DEFAULT_INVALID_OUTPUT)
    parser.add_argument("--summary-output", type=Path, default=DEFAULT_SUMMARY_OUTPUT)
    if DATABASE == "30_CircR2Disease":
        parser.add_argument("--preprocessed-output", type=Path, default=SCRIPT_DIR / "clean" / f"{DATABASE}.entries.mouse.preprocessed.tsv")
    parser.add_argument("--force", action="store_true")
    return parser.parse_args()


def norm(value: Any) -> str:
    if value is None:
        return "NA"
    text = str(value).strip()
    return "NA" if text in NULL_VALUES or text.upper() in NULL_VALUES else text


def norm_strand(value: Any) -> str:
    if value is None:
        return "NA"
    text = str(value).strip()
    return text if text in {"+", "-"} else "NA"


def norm_chrom(value: Any) -> str:
    text = norm(value)
    if text == "NA" or text.startswith("chr"):
        return text
    if re.fullmatch(r"[0-9XYM]+", text):
        return f"chr{text}"
    return text


def int_text(value: Any) -> str:
    text = norm(value)
    if text == "NA":
        raise ValueError("missing coordinate")
    return str(int(float(text)))


def parse_chr_pipe(value: Any) -> tuple[str, str, str] | None:
    match = re.match(r"^(?P<chrom>[^:]+):(?P<start>\d+)\|(?P<end>\d+)$", norm(value))
    if not match:
        return None
    return norm_chrom(match.group("chrom")), match.group("start"), match.group("end")


def parse_chr_dash(value: Any) -> tuple[str, str, str] | None:
    match = re.match(r"^(?P<chrom>[^:]+):(?P<start>\d+)-(?P<end>\d+)$", norm(value))
    if not match:
        return None
    return norm_chrom(match.group("chrom")), match.group("start"), match.group("end")


def parse_position_strand(value: Any) -> tuple[str, str, str, str] | None:
    match = re.match(r"^(?P<chrom>[^:]+):(?P<start>\d+)-(?P<end>\d+)\|(?P<strand>[+-])$", norm(value))
    if not match:
        return None
    return norm_chrom(match.group("chrom")), match.group("start"), match.group("end"), match.group("strand")


def validate_row(row: dict[str, str]) -> None:
    if row["circRNA_id"] == "NA" or row["chrom"] == "NA":
        raise ValueError("missing circRNA_id or chrom")
    if int(row["start"]) > int(row["end"]):
        raise ValueError("start greater than end")
    if row["strand"] not in {"+", "-", "NA"}:
        raise ValueError("invalid strand")


def read_dict_tsv(path: Path) -> Iterable[dict[str, str]]:
    with path.open(newline="") as handle:
        yield from csv.DictReader(handle, delimiter="\t")


def read_dict_csv(path: Path) -> Iterable[dict[str, str]]:
    with path.open(newline="") as handle:
        yield from csv.DictReader(handle)


def dedup_rows(rows: Iterable[dict[str, str]], counts: Counter[str]) -> list[dict[str, str]]:
    seen = set()
    out = []
    for row in rows:
        key = tuple(row[column] for column in OUTPUT_COLUMNS)
        if key in seen:
            counts["duplicate_full_rows_removed"] += 1
            continue
        seen.add(key)
        out.append(row)
    return out


def write_tsv(path: Path, rows: list[dict[str, str]], columns: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, delimiter="\t", lineterminator="\n", extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_summary(path: Path, counts: Counter[str], extras: dict[str, str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(["metric", "value"])
        for key in sorted(counts):
            writer.writerow([key, counts[key]])
        for key, value in sorted(extras.items()):
            writer.writerow([key, value])


def refuse_outputs(paths: list[Path], force: bool) -> None:
    if force:
        return
    existing = [str(path) for path in paths if path.exists()]
    if existing:
        raise FileExistsError("Output exists; use --force: " + ", ".join(existing))

def normalize_region_text(value: Any) -> str:
    text = unicodedata.normalize("NFKC", norm(value))
    for ch in ["\u2010", "\u2011", "\u2012", "\u2013", "\u2014", "\u2015", "\u2212", "\ufe58", "\ufe63", "\uff0d"]:
        text = text.replace(ch, "-")
    return re.sub(r"\s+", "", text)


def clean_rows(raw_dir: Path, counts: Counter[str], invalid: list[dict[str, str]], args: argparse.Namespace) -> list[dict[str, str]]:
    if load_workbook is None:
        raise RuntimeError("openpyxl is required for CircR2Disease clean")
    path = raw_dir / "The circRNA-disease entries.xlsx"
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    iterator = ws.iter_rows(values_only=True)
    headers = [str(v) for v in next(iterator)]
    rows, preprocessed = [], []
    for line_number, values in enumerate(iterator, start=2):
        raw = {header: values[i] if i < len(values) else None for i, header in enumerate(headers)}
        counts["raw_rows"] += 1
        if norm(raw.get("Species")) not in {"Mouse", "Mus musculus"}:
            counts["non_mouse_rows"] += 1
            continue
        counts["mouse_rows"] += 1
        pre = {key: norm(value) for key, value in raw.items()}
        pre["circRNA_id"] = norm(raw.get("circRNA Name"))
        pre["region_normalized"] = normalize_region_text(raw.get("Region"))
        pre["strand_normalized"] = norm_strand(raw.get("Strand"))
        preprocessed.append(pre)
        try:
            parsed = parse_chr_dash(pre["region_normalized"])
            if not parsed:
                raise ValueError("invalid Region")
            chrom, start, end = parsed
            clean = {"circRNA_id": pre["circRNA_id"], "chrom": chrom, "start": start, "end": end, "strand": pre["strand_normalized"], "gene_symbol": norm(raw.get("Gene Symbol"))}
            validate_row(clean)
            rows.append(clean)
        except Exception as exc:
            invalid.append({"source_file": path.name, "line_number": str(line_number), "invalid_reason": str(exc), **pre})
            counts["invalid_rows"] += 1
    pre_cols = list(preprocessed[0].keys()) if preprocessed else headers + ["circRNA_id", "region_normalized", "strand_normalized"]
    write_tsv(args.preprocessed_output, preprocessed, pre_cols)
    counts["preprocessed_rows"] = len(preprocessed)
    return rows

def main() -> int:
    args = parse_args()
    output_paths = [args.output, args.invalid_output, args.summary_output]
    if hasattr(args, "preprocessed_output"):
        output_paths.append(args.preprocessed_output)
    refuse_outputs(output_paths, args.force)
    counts: Counter[str] = Counter()
    invalid: list[dict[str, str]] = []
    rows = clean_rows(args.raw_dir, counts, invalid, args)
    counts["clean_rows_before_uniq"] = len(rows)
    rows = dedup_rows(rows, counts)
    counts["clean_rows"] = len(rows)
    counts["unique_circRNA_ids"] = len({row["circRNA_id"] for row in rows})
    counts["stranded_rows"] = sum(1 for row in rows if row["strand"] in {"+", "-"})
    counts["unstranded_rows"] = sum(1 for row in rows if row["strand"] == "NA")
    write_tsv(args.output, rows, OUTPUT_COLUMNS)
    invalid_columns = sorted({key for row in invalid for key in row}) or ["source_file", "line_number", "invalid_reason"]
    write_tsv(args.invalid_output, invalid, invalid_columns)
    extras = {"raw_dir": str(args.raw_dir), "output": str(args.output), "invalid_output": str(args.invalid_output)}
    if hasattr(args, "preprocessed_output"):
        extras["preprocessed_output"] = str(args.preprocessed_output)
    write_summary(args.summary_output, counts, extras)
    print(f"Wrote {args.output}")
    print(f"Wrote {args.invalid_output}")
    print(f"Wrote {args.summary_output}")
    print(f"Summary: {counts['raw_rows']} raw rows, {counts['clean_rows']} clean rows, {counts['invalid_rows']} invalid rows")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
